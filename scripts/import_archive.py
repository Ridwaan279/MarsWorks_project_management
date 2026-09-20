"""
One-off migration: read the sub-teams' existing planners out of archive/ and
emit prisma/seed-data.json in the canonical shape.

This is deliberately a throwaway tool for the .xlsx files the teams used
before consolidating. The live Google Sheets sync will use the Sheets API from
the Node side, not this. What is worth keeping here is the *column mapping* --
which field in each team's planner becomes which canonical field -- because
that mapping is the thing the sync has to agree with.

    pip install openpyxl && python3 scripts/import_archive.py
"""

import json
import re
from datetime import datetime
from pathlib import Path

import openpyxl

ARCHIVE = Path(__file__).resolve().parent.parent / "archive"
OUT = Path(__file__).resolve().parent.parent / "prisma" / "seed-data.json"

# Status vocabularies across the three planners, mapped onto ours. Mirrors
# STATUS_ALIASES in src/lib/domain.ts -- keep the two in step.
STATUS = {
    "idea": "BACKLOG",
    "backlog": "BACKLOG",
    "not started": "TODO",
    "to do": "TODO",
    "todo": "TODO",
    "in progress": "IN_PROGRESS",
    "in-progress": "IN_PROGRESS",
    "blocked": "BLOCKED",
    "in review": "IN_REVIEW",
    "done": "DONE",
    "complete": "DONE",
    "completed": "DONE",
}

STAGE = {
    "1": "INVESTIGATION",
    "2": "DESIGN",
    "3": "PROTOTYPE",
    "4": "ORDER",
    "5": "CONSTRUCTION",
    "6": "TESTING",
    "7": "IMPROVEMENT",
    "8": "DOCUMENTATION",
}

# Master-timeline tab name -> canonical team key.
#
# There are six sub-teams. The master timeline also has a "Leadership and
# Milestones" tab and tabs for Drone and Mini-Rover. The milestone rows on the
# leadership tab become Milestone records; its remaining rows are recruitment
# and administration, which belong to Operations. The Drone and Mini-Rover tabs
# hold no tasks and no sub-team owns them, so they are not imported.
TAB_TO_TEAM = {
    "Leadership and Milestones": "OPS",
    "Elec Tasks": "ELEC",
    "Mech Tasks": "MECH",
    "Robotics Tasks": "ROBO",
    "Sci Tasks": "SCI",
    "Software Tasks": "SW",
}


def norm_status(raw):
    if raw is None:
        return None
    return STATUS.get(str(raw).strip().lower())


def iso(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    return None


def clean(value):
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def import_master():
    """Master Google Sheet: Tasks | Assignee | Start | End | Status | Notes.

    Rows whose status is literally "Milestone" are not tasks -- they are dated
    checkpoints with start == end, so they become Milestone records."""
    wb = openpyxl.load_workbook(ARCHIVE / "Master Timeline - Project MarsWorks 2026.xlsx", data_only=True)
    tasks, milestones = [], []
    for ws in wb.worksheets:
        team = TAB_TO_TEAM.get(ws.title)
        if not team:
            continue
        for row in list(ws.iter_rows(values_only=True))[1:]:
            title = clean(row[0])
            if not title:
                continue
            owner, start, end, status, notes = (
                clean(row[1]), iso(row[2]), iso(row[3]), clean(row[4]), clean(row[5]),
            )
            if status and status.lower() == "milestone":
                milestones.append({
                    "name": re.sub(r"^Milestone:\s*", "", title),
                    "targetDate": start or end,
                    "description": notes,
                })
                continue
            tasks.append({
                "team": team, "title": title, "ownerLabel": owner,
                "plannedStart": start, "plannedEnd": end,
                "status": norm_status(status) or "TODO",
                "notes": notes, "source": "master-sheet",
            })
    return tasks, milestones


def import_mechanical():
    """Smartsheet Gantt: WBS NUMBER | TASK TITLE | TASK OWNER | START | DUE |
    DURATION | PCT COMPLETE. An x.0 row is a workstream heading; x.y rows are
    its tasks."""
    wb = openpyxl.load_workbook(ARCHIVE / "Mechanical_Gantt chart.xlsx", data_only=True)
    ws = wb["Gantt Chart"]
    workstreams, tasks = [], []
    current = None
    for row in ws.iter_rows(min_row=11, values_only=True):
        wbs, title, owner = clean(row[1]), clean(row[2]), clean(row[3])
        if not wbs or not title:
            continue
        if wbs.endswith(".0"):
            current = title
            workstreams.append({"team": "MECH", "code": wbs, "name": title})
            continue
        pct = row[7]
        tasks.append({
            "team": "MECH", "workstream": current, "code": wbs, "title": title,
            "ownerLabel": owner, "plannedStart": iso(row[4]), "plannedEnd": iso(row[5]),
            "progress": int(round(float(pct) * 100)) if isinstance(pct, (int, float)) else 0,
            "status": None, "source": "mech-gantt",
        })
    return workstreams, tasks


def import_electrical():
    """Electrical dashboard: Project Stage | Task Title | Description |
    Assigned To | Status | Start | Due | Critical? | Source Link | ...

    A row with a Description but no Task Title is a checklist item belonging to
    the task above it, not a task of its own."""
    wb = openpyxl.load_workbook(ARCHIVE / "Electrical Project Dashboard.xlsx", data_only=True)
    ws = wb["TASK tracker"]
    tasks = []
    for row in list(ws.iter_rows(values_only=True))[1:]:
        stage_raw, title, desc = clean(row[0]), clean(row[1]), clean(row[2])
        if title:
            tasks.append({
                "team": "ELEC",
                "workstream": re.sub(r"^\d+\.\s*", "", stage_raw) if stage_raw else None,
                "title": title, "description": desc, "ownerLabel": clean(row[3]),
                "status": norm_status(row[4]) or "TODO",
                "plannedStart": iso(row[5]), "plannedEnd": iso(row[6]),
                "stage": STAGE.get(stage_raw.split(".")[0]) if stage_raw else None,
                "priority": "CRITICAL" if row[7] is True else "MEDIUM",
                "link": clean(row[8]), "subtasks": [], "source": "elec-dashboard",
            })
        elif desc and tasks:
            # Checklist row: belongs to the task above.
            tasks[-1]["subtasks"].append(desc)
    return tasks


def main():
    master_tasks, milestones = import_master()
    mech_ws, mech_tasks = import_mechanical()
    elec_tasks = import_electrical()

    # The master sheet's Mech and Elec tabs are empty; those teams keep their
    # plans in their own files. Drop master rows for teams with a richer source
    # so nothing is counted twice.
    richer = {"MECH", "ELEC"}
    master_tasks = [t for t in master_tasks if t["team"] not in richer]

    payload = {
        "workstreams": mech_ws,
        "milestones": milestones,
        "tasks": master_tasks + mech_tasks + elec_tasks,
    }
    OUT.write_text(json.dumps(payload, indent=2))
    by_team = {}
    for t in payload["tasks"]:
        by_team[t["team"]] = by_team.get(t["team"], 0) + 1
    print(f"wrote {OUT.relative_to(OUT.parent.parent)}")
    print(f"  milestones : {len(payload['milestones'])}")
    print(f"  workstreams: {len(payload['workstreams'])}")
    print(f"  tasks      : {len(payload['tasks'])}  {by_team}")
    print(f"  subtasks   : {sum(len(t.get('subtasks', [])) for t in payload['tasks'])}")


if __name__ == "__main__":
    main()
